"""Tests du Module 12 — pytest export-terrain/ -v (aucune base requise)."""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from export_excel import construire_classeur


def _donnees_synthetiques() -> dict:
    return {
        "prix": [
            ("ACIER_T", "Acier", "tonne", date(2026, 7, 1), 950.0, 8.6, True),
            ("BETON_M3", "Béton", "m3", date(2026, 7, 1), 130.0, 0.6, False),
        ],
        "appels_offres": [
            (date(2026, 9, 1), date(2026, 10, 1), "69", 7,
             "Travaux de maçonnerie", "Mairie X", "https://boamp.fr/x"),
        ],
        "devis": [
            (date(2026, 1, 1), "Client A", "Maison individuelle", "01", 120.0,
             250000.0, "gagné", 8.5),
            (date(2026, 2, 1), "Client B", "VRD / voirie", "69", 5000.0,
             400000.0, "perdu", 4.0),
        ],
    }


def test_classeur_a_quatre_feuilles_avec_les_bons_totaux(tmp_path):
    wb = construire_classeur(_donnees_synthetiques())
    chemin = tmp_path / "test.xlsx"
    wb.save(chemin)

    relu = load_workbook(chemin)
    assert relu.sheetnames == ["Synthèse", "Prix matériaux", "Appels d'offres", "Devis"]

    feuille_prix = relu["Prix matériaux"]
    assert feuille_prix.max_row == 3  # 1 entête + 2 lignes
    assert feuille_prix["A2"].value == "ACIER_T"

    feuille_devis = relu["Devis"]
    assert feuille_devis.max_row == 3
    assert feuille_devis["G2"].value == "gagné"


def test_synthese_contient_le_taux_de_reussite():
    wb = construire_classeur(_donnees_synthetiques())
    ws = wb["Synthèse"]
    valeurs = [ws.cell(row=r, column=1).value for r in range(1, 10)]
    assert "Taux de réussite des devis" in valeurs
