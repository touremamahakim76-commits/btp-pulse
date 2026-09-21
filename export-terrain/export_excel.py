"""Module 12 — Export Excel pour les équipes terrain (BTP Pulse).

Certaines équipes n'ont pas accès à Power BI : ce script génère un
classeur Excel autonome (4 feuilles) avec mise en forme, pour une
consultation directe sur le terrain (tablette, PC chantier) ou un envoi
par mail.

Usage :
  python export_excel.py
  python export_excel.py --sortie mon_export.xlsx
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "database"))
import db_config  # noqa: E402

DOSSIER_EXPORTS = Path(__file__).resolve().parent / "exports"
DOSSIER_EXPORTS.mkdir(exist_ok=True)

ENTETE_FOND = PatternFill("solid", fgColor="1E3A5F")
ENTETE_POLICE = Font(color="FFFFFF", bold=True)
TITRE_POLICE = Font(bold=True, size=14, color="1E3A5F")
VERT = PatternFill("solid", fgColor="E3F5EC")
ROUGE = PatternFill("solid", fgColor="FBE7E7")


# =========================================================================
# Lecture des données
# =========================================================================
def lire_donnees(conn) -> dict:
    with conn.cursor() as cur:
        cur.execute("""
            SELECT v.code_materiau, v.libelle, v.unite, v.date_releve, v.prix,
                   t.variation_pct, t.significatif
            FROM v_dernier_prix_materiau v
            LEFT JOIN LATERAL (
                SELECT variation_pct, significatif FROM resultats_tests_statistiques
                WHERE code_materiau = v.code_materiau ORDER BY date_test DESC LIMIT 1
            ) t ON true
            ORDER BY v.code_materiau
        """)
        prix = cur.fetchall()

        cur.execute("""
            SELECT date_parution, date_limite_reponse, departements,
                   score_pertinence, objet, acheteur, url_avis
            FROM appels_offres
            WHERE score_pertinence >= 3
              AND (date_limite_reponse IS NULL OR date_limite_reponse >= current_date)
            ORDER BY score_pertinence DESC, date_parution DESC
            LIMIT 200
        """)
        appels_offres = cur.fetchall()

        cur.execute("""
            SELECT date_devis, client, type_chantier, departement, surface_m2,
                   montant_ht, statut, marge_estimee_pct
            FROM devis_historique ORDER BY date_devis DESC LIMIT 300
        """)
        devis = cur.fetchall()

    return {"prix": prix, "appels_offres": appels_offres, "devis": devis}


# =========================================================================
# Mise en forme générique
# =========================================================================
def _ecrire_feuille(ws, entetes: list[str], lignes: list[tuple]) -> None:
    ws.append(entetes)
    for cell in ws[1]:
        cell.fill = ENTETE_FOND
        cell.font = ENTETE_POLICE
        cell.alignment = Alignment(vertical="center")
    for ligne in lignes:
        ws.append(list(ligne))
    ws.freeze_panes = "A2"
    for i, entete in enumerate(entetes, start=1):
        largeur = max(12, min(45, len(str(entete)) + 4))
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.auto_filter.ref = ws.dimensions


# =========================================================================
# Feuilles
# =========================================================================
def feuille_synthese(wb: Workbook, donnees: dict) -> None:
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = "BTP Pulse — export terrain"
    ws["A1"].font = TITRE_POLICE
    ws["A2"] = f"Généré le {datetime.now():%d/%m/%Y à %H:%M}"

    nb_alertes = sum(1 for p in donnees["prix"] if p[6])
    nb_ao = len(donnees["appels_offres"])
    devis = donnees["devis"]
    gagnes = sum(1 for d in devis if d[6] == "gagné")
    statues = sum(1 for d in devis if d[6] in ("gagné", "perdu"))
    taux = round(100 * gagnes / statues, 1) if statues else 0

    lignes = [
        ("Matériaux suivis", len(donnees["prix"])),
        ("Variations de prix significatives", nb_alertes),
        ("Appels d'offres pertinents ouverts", nb_ao),
        ("Devis historiques", len(devis)),
        ("Taux de réussite des devis", f"{taux} %"),
    ]
    for i, (label, valeur) in enumerate(lignes, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valeur)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18


def feuille_prix(wb: Workbook, donnees: dict) -> None:
    ws = wb.create_sheet("Prix matériaux")
    entetes = ["Code", "Libellé", "Unité", "Dernier relevé", "Prix (€)",
              "Variation (%)", "Significatif ?"]
    lignes = [
        (p[0], p[1], p[2], p[3], float(p[4]),
         float(p[5]) if p[5] is not None else None,
         "Oui" if p[6] else "Non")
        for p in donnees["prix"]
    ]
    _ecrire_feuille(ws, entetes, lignes)
    n = len(lignes)
    if n:
        ws.conditional_formatting.add(
            f"F2:F{n + 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=ROUGE)
        )
        ws.conditional_formatting.add(
            f"F2:F{n + 1}", CellIsRule(operator="lessThan", formula=["0"], fill=VERT)
        )


def feuille_appels_offres(wb: Workbook, donnees: dict) -> None:
    ws = wb.create_sheet("Appels d'offres")
    entetes = ["Parution", "Limite réponse", "Départements", "Score",
              "Objet", "Acheteur", "Lien avis"]
    _ecrire_feuille(ws, entetes, donnees["appels_offres"])


def feuille_devis(wb: Workbook, donnees: dict) -> None:
    ws = wb.create_sheet("Devis")
    entetes = ["Date", "Client", "Type de chantier", "Département",
              "Surface (m²)", "Montant HT (€)", "Statut", "Marge estimée (%)"]
    lignes = [
        (d[0], d[1], d[2], d[3], float(d[4]) if d[4] else None,
         float(d[5]), d[6], float(d[7]) if d[7] else None)
        for d in donnees["devis"]
    ]
    _ecrire_feuille(ws, entetes, lignes)
    n = len(lignes)
    if n:
        ws.conditional_formatting.add(
            f"G2:G{n + 1}",
            CellIsRule(operator="equal", formula=['"gagné"'], fill=VERT),
        )
        ws.conditional_formatting.add(
            f"G2:G{n + 1}",
            CellIsRule(operator="equal", formula=['"perdu"'], fill=ROUGE),
        )


# =========================================================================
# CLI
# =========================================================================
def construire_classeur(donnees: dict) -> Workbook:
    wb = Workbook()
    feuille_synthese(wb, donnees)
    feuille_prix(wb, donnees)
    feuille_appels_offres(wb, donnees)
    feuille_devis(wb, donnees)
    return wb


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Export Excel terrain — BTP Pulse")
    p.add_argument("--sortie", default=None, help="chemin du fichier .xlsx généré")
    args = p.parse_args(argv)

    for _flux in (sys.stdout, sys.stderr):
        try:
            _flux.reconfigure(encoding="utf-8")
        except Exception:
            pass

    try:
        conn = psycopg2.connect(db_config.dsn())
    except psycopg2.OperationalError as exc:
        print(f"ERREUR : connexion PostgreSQL impossible.\n{exc}")
        return 1

    try:
        donnees = lire_donnees(conn)
    finally:
        conn.close()

    wb = construire_classeur(donnees)
    sortie = Path(args.sortie) if args.sortie else DOSSIER_EXPORTS / f"btp_pulse_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(sortie)
    print(f"Export généré : {sortie}")
    print(f"  {len(donnees['prix'])} matériaux, {len(donnees['appels_offres'])} appels d'offres, "
          f"{len(donnees['devis'])} devis")
    return 0


if __name__ == "__main__":
    sys.exit(main())
